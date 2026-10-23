"""ATS-friendly DOCX: built-in Heading/List styles, no tables, text boxes or header/footer content."""

from io import BytesIO
from datetime import datetime, timezone
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.export.content import Layout

DARK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x44, 0x44, 0x44)


def _bottom_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": color}.items():
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _para(doc, text: str, size: float = 10, bold=False, italic=False, color=DARK, align=None, after=0, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.size, run.bold, run.italic, run.font.color.rgb = Pt(size), bold, italic, color
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def render_docx(layout: Layout) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    doc.core_properties.title = f"CV - {layout.name}"
    doc.core_properties.author = layout.name
    doc.core_properties.created = doc.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    doc.core_properties.last_modified_by = "PasseFiltre"
    doc.core_properties.revision = 1

    center = WD_ALIGN_PARAGRAPH.CENTER
    accent = RGBColor.from_string(layout.accent.lstrip("#").upper())
    _para(doc, layout.name, 20, bold=True, color=accent, align=center)
    if layout.headline:
        _para(doc, layout.headline, 11, color=GREY, align=center)
    if layout.contact:
        _para(doc, " | ".join(layout.contact), 9, color=GREY, align=center, after=4)

    for section in layout.sections:
        h = doc.add_heading(level=1)
        run = h.add_run(section.title.upper())
        run.font.size, run.font.color.rgb, run.font.name = Pt(11.5), accent, "Calibri"
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(10), Pt(4)
        _bottom_border(h, layout.accent.lstrip("#").upper())
        if section.text:
            _para(doc, section.text)
        for block in section.blocks:
            title = _para(doc, block.title, 10.5, bold=True)
            title.paragraph_format.space_before = Pt(5)
            title.paragraph_format.keep_with_next = True
            if block.meta:
                _para(doc, block.meta, 9.5, italic=True, color=GREY).paragraph_format.keep_with_next = True
            for t in block.paragraphs:
                _para(doc, t)
            for t in block.bullets:
                _para(doc, t, style="List Bullet")
            for t in block.footer:
                _para(doc, t, 9.5, italic=True, color=GREY)

    buf = BytesIO()
    doc.save(buf)
    # OPC ZIP timestamps otherwise depend on wall clock, even with fixed document metadata.
    stable = BytesIO()
    with ZipFile(buf) as source, ZipFile(stable, "w") as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return stable.getvalue()
